"""DP_2026 Demand Planning çalışma kitabını planlama ekranının veri dosyasına çevirir.

Kullanım:
    python import_dp2026.py <Demand_Planning.xlsx> <cikti.json>

Okunan sayfalar: 2_PAZ_PLAN, 3_SAT_PLAN, 4_SELL_IN, 5_SELL_OUT, 1_HEDEFLER, ÖZET,
LISTE_FIYATLARI, MUSTERI_KOSULLARI. Değerler Excel'in son kaydedilmiş (hesaplanmış) halinden okunur.
Tüm ciro değerleri dosyadaki gibi KDV hariçtir.
"""
import json
import sys
from collections import Counter, defaultdict

import openpyxl

MONTHS = ['Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim', 'Kasım', 'Aralık']
# Sell-in / sell-out dosyasındaki ticari unvan -> plan dosyasındaki kısa müşteri adı
CUSTOMER = {
    'MEDİA MARKT TURKEY TİC.LTD.ŞTİ.': 'Media Markt', 'VATAN BİLGİSAYAR SAN.TİC.A.Ş.': 'Vatan',
    'TEKNOSA İÇ VE DIŞ TİC.A.Ş.': 'Teknosa', 'AMAZON TURKEY PERAKENDE HİZMETLERİ LİMİTED ŞİRKETİ': 'Amazon',
    'QUAD TEKNOLOJİ TİCARET A.Ş.': 'QUAD', 'EVKUR ALIŞVERİŞ MERKEZLERİ TİCARET A.Ş.': 'Evkur',
    'ARÇELİK PAZARLAMA ANONİM ŞİRKETİ': 'Arçelik', 'YENİ MAĞAZACILIK ANONİM ŞİRKETİ': 'A-101', 'BLU': 'BLU',
    'NİT-SET EV ALETLERİ PAZARLAMA SANAYİ VE TİCARET LİMİTED ŞİRKETİ': 'Dist. TV',
    'Gİ-PA DAYANIKLI TÜKETİM MAMÜLLERİ TİCARET ANONİM ŞİRKETİ': 'Gİ-PA',
    'SAHİP DAYANIKLI TÜKETİM MALLARI İMALAT VE PAZARLAMA TİC.LTD.ŞTİ.': 'SAHİP',
}
TARGET_CUSTOMER = {'Yeni Mağazacılık': 'A-101', 'NİT-SET': 'Dist. TV'}
KOSUL_COLS = {5: 'Vatan', 6: 'Media Markt', 7: 'Teknosa', 8: 'QUAD', 9: 'Diğer', 10: 'Amazon', 11: 'Evkur', 12: 'Arçelik'}


def code(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def n(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def r2(v):
    return int(v) if float(v).is_integer() else round(v, 2)


def trim(a):
    a = [r2(x) for x in a]
    while a and a[-1] == 0:
        a.pop()
    return a


def main(src, dst):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = lambda name, start: (r for r in wb[name].iter_rows(min_row=start, values_only=True) if r)
    issues = Counter()
    P = {}

    def prod(c):
        return P.setdefault(c, {'k': c, 'cu': {}})

    for r in rows('2_PAZ_PLAN', 4):
        c = code(r[0])
        if not c:
            continue
        p = prod(c)
        p.update(d=r[1] or '', b=r[2] or 'Bilinmiyor', g=r[3] or 'Grupsuz', s=r[4] or '', h=1 if n(r[5]) else 0,
                 a=r[6] if isinstance(r[6], (int, float)) else None, st=n(r[7]), it=n(r[8]),
                 pz=trim([n(x) for x in r[9:24]]), pzc=trim([n(x) for x in r[25:40]]), al=trim([n(x) for x in r[41:56]]))
        if not r[2]:
            issues['PAZ_PLAN: markası boş ürün'] += 1

    for r in rows('3_SAT_PLAN', 4):
        c = code(r[1])
        if not c or not r[0]:
            continue
        q = [n(x) for x in r[14:26]] + [n(x) for x in r[53:56]]
        qc = [n(x) for x in r[27:39]] + [n(x) for x in r[56:59]]
        p = prod(c)
        if 'd' not in p:
            p.update(d=r[2] or '', b=r[3] or 'Bilinmiyor', g=r[4] or 'Grupsuz', s=r[5] or '', h=1 if n(r[6]) else 0,
                     a=None, st=0, it=0, pz=[], pzc=[], al=[])
            issues['SAT_PLAN\'da olup PAZ_PLAN\'da olmayan ürün'] += 1
        if any(q) or n(r[12]) or n(r[13]):
            cu = p['cu'].setdefault(r[0], {})
            cu.update(sp=trim(q), spc=trim(qc), m=r2(n(r[40])), sio=r2(n(r[12])), soo=r2(n(r[13])))

    for sheet, key in (('4_SELL_IN', 'si'), ('5_SELL_OUT', 'so')):
        unknown = Counter()
        for r in rows(sheet, 4):
            c = code(r[1])
            if not c or r[4] not in MONTHS:
                continue
            cust = CUSTOMER.get(r[0])
            if not cust:
                unknown[r[0]] += 1
                cust = str(r[0])
            p = prod(c)
            if 'd' not in p:
                p.update(d='', b=r[3] or 'Bilinmiyor', g=r[2] or 'Grupsuz', s='', h=0, a=None, st=0, it=0, pz=[], pzc=[], al=[])
                issues[f'{sheet}: plan dosyasında olmayan ürün kodu'] += 1
            cu = p['cu'].setdefault(cust, {})
            mi = MONTHS.index(r[4])
            arr = cu.setdefault(key, [0] * 12)
            arrc = cu.setdefault(key + 'c', [0] * 12)
            arr[mi] += n(r[5])
            arrc[mi] += n(r[6])
            if n(r[5]) < 0:
                issues[f'{sheet}: negatif adet (iade?)'] += 1
        for k, v in unknown.items():
            issues[f'{sheet}: eşleşmeyen müşteri adı "{k}"'] += v

    for r in rows('ÖZET', 4):
        c = code(r[0])
        if not c or c not in P:
            continue
        p = P[c]
        p.update(kn=r[4] or '', y=[x for x in (('13 Ağustos', r[5]), ('8 Ekim', r[6])) if x[1]],
                 sv=r2(n(r[15])), srf=r2(n(r[14])) if isinstance(r[14], (int, float)) else None)

    hdr = None
    for r in wb['LISTE_FIYATLARI'].iter_rows(min_row=1, values_only=True):
        if hdr is None:
            hdr = r
            continue
        c = code(r[0])
        if c in P:
            P[c]['tk'] = trim([n(r[4 + 2 * i]) for i in range(15)])
            P[c]['pr'] = trim([n(r[5 + 2 * i]) for i in range(15)])

    for r in rows('MUSTERI_KOSULLARI', 2):
        c = code(r[1])
        if c in P:
            P[c]['ko'] = {name: r2(n(r[i])) for i, name in KOSUL_COLS.items() if n(r[i])}

    targets = []
    for r in rows('1_HEDEFLER', 4):
        if not r[0] or r[2] in ('UGREEN TOPLAM', 'PHILIPS TOPLAM', 'MÜŞTERİ TOPLAM') or r[1] == 'TÜM MARKA':
            continue
        vals = [n(x) for x in r[3:15]]
        if any(vals):
            targets.append([TARGET_CUSTOMER.get(r[0], r[0]), r[1], r[2], [r2(v) for v in vals]])

    for p in P.values():
        for cu in p['cu'].values():
            for k in ('si', 'sic', 'so', 'soc'):
                if k in cu:
                    cu[k] = trim(cu[k])

    out = {'kaynak': 'DP_2026_v2 Demand Planning', 'urunler': list(P.values()), 'hedefler': targets,
           'sorunlar': dict(issues)}
    with open(dst, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, separators=(',', ':'))
    print(f'{len(P)} ürün, {len(targets)} hedef satırı; sorunlar: {dict(issues)}')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
